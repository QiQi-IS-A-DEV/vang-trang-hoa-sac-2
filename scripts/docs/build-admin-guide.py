"""Build the editable admin handbook and its Markdown companion."""
import json
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'docs/user-guide'
pages = json.loads((OUT / 'content.json').read_text(encoding='utf-8'))
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(.65)
sec.left_margin = sec.right_margin = Inches(.75)
sec.header_distance = sec.footer_distance = Inches(.25)
for name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'Caption', 'List Bullet', 'List Number']:
    s = doc.styles[name]
    s.font.name = 'Arial'
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.font.size = Pt(11)
    s.paragraph_format.space_after = Pt(6)
    s.paragraph_format.line_spacing = 1.08
doc.styles['Title'].font.size = Pt(26)
doc.styles['Heading 1'].font.size = Pt(20)
doc.styles['Heading 2'].font.size = Pt(13)
doc.styles['Caption'].font.size = Pt(9)
doc.styles['Caption'].font.italic = True
doc.styles['Heading 1'].paragraph_format.space_after = Pt(10)
hp = sec.header.paragraphs[0]
hp.text = 'VẦNG TRĂNG HÒA SẮC 2   •   SỔ TAY QUẢN TRỊ'
hp.runs[0].font.size = Pt(8)
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
fp.add_run('Hướng dẫn quản trị  •  ' ).font.size = Pt(8)
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE'); fp._p.append(fld)

def hyperlink(p, label, url):
    rid = p.part.relate_to(url, 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True)
    el = OxmlElement('w:hyperlink'); el.set(qn('r:id'), rid)
    r = OxmlElement('w:r'); prop = OxmlElement('w:rPr')
    u = OxmlElement('w:u'); u.set(qn('w:val'), 'single'); prop.append(u); r.append(prop)
    t = OxmlElement('w:t'); t.text = label; r.append(t); el.append(r); p._p.append(el)

md = ['# Hướng dẫn quản trị Vầng Trăng Hòa Sắc 2', '', 'Cập nhật 05/10/2026. Tài liệu không chứa mật khẩu.', '']
for index, page in enumerate(pages):
    heading = doc.add_paragraph(page['title'], 'Title' if index == 0 else 'Heading 1')
    if index: heading.paragraph_format.page_break_before = True
    md += ['## ' + page['title'], '']
    if page.get('route'):
        doc.add_paragraph('Mở màn hình: ' + page['route'])
        md += ['**Mở màn hình:** `' + page['route'] + '`', '']
    for kind, value in page['blocks']:
        if kind == 'p':
            doc.add_paragraph(value); md += [value, '']
        elif kind == 'h':
            doc.add_paragraph(value, 'Heading 2'); md += ['### ' + value, '']
        elif kind in ('steps', 'bullets'):
            for n, text in enumerate(value, 1):
                # Explicit numbers restart each short workflow reliably across Word and LibreOffice.
                p = doc.add_paragraph((str(n) + '. ' if kind == 'steps' else '• ') + text)
                p.paragraph_format.left_indent = Inches(.17)
                p.paragraph_format.first_line_indent = Inches(-.17)
                md.append((str(n)+'. ' if kind == 'steps' else '- ') + text)
            md.append('')
        elif kind == 'table':
            table = doc.add_table(rows=0, cols=len(value[0])); table.style = 'Light Shading Accent 1'
            for row_index, data in enumerate(value):
                row = table.add_row()
                for cell, text in zip(row.cells, data):
                    cell.text = str(text)
                    for p in cell.paragraphs:
                        p.paragraph_format.space_after = Pt(4)
                        p.paragraph_format.space_before = Pt(4)
                        for r in p.runs: r.font.size = Pt(10.5); r.bold = row_index == 0
                trpr = row._tr.get_or_add_trPr()
                nosplit = OxmlElement('w:cantSplit'); trpr.append(nosplit)
                if row_index == 0:
                    repeat = OxmlElement('w:tblHeader'); trpr.append(repeat)
            md += ['| ' + ' | '.join(map(str,value[0])) + ' |', '| ' + ' | '.join(['---']*len(value[0])) + ' |']
            md += ['| ' + ' | '.join(str(x).replace('|','/') for x in row) + ' |' for row in value[1:]] + ['']
        elif kind in ('image', 'logo'):
            file, caption = (OUT / 'screenshots' / value[0], value[1]) if kind == 'image' else (OUT / value, '')
            # Keep screenshots at readable width; no stretching, cropping, or resampling.
            width = 6.8 if kind == 'image' else 1.35
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            shape = p.add_run().add_picture(str(file), width=Inches(width))
            shape._inline.docPr.set('descr', caption or 'Logo chương trình Vầng Trăng Hòa Sắc 2')
            if caption: doc.add_paragraph(caption, 'Caption')
            md += ['![' + (caption or 'Logo chương trình') + '](' + ('screenshots/'+value[0] if kind=='image' else value) + ')', '']
        elif kind == 'links':
            for label, url in value:
                hyperlink(doc.add_paragraph(), label + ': ' + url, url)
                md += ['['+label+']('+url+')', '']
doc.core_properties.title = pages[0]['title']
doc.core_properties.subject = 'Hướng dẫn tất cả tính năng quản trị dành cho BTC'
doc.core_properties.author = 'Vầng Trăng Hòa Sắc 2'
doc.save(OUT / 'Huong-dan-quan-tri-VTHS2.docx')
(OUT / 'HUONG-DAN-QUAN-TRI.md').write_text('\n'.join(md), encoding='utf-8')
print('Built DOCX and Markdown')
